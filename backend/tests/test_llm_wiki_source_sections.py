"""Source structure, stable relations and source-scoped assignment regressions."""

from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace

import pytest

from backend.domains.configuration.source_sections_schema import ensure_source_sections
from backend.domains.llm_wiki import documents, origins, planning, writing
from backend.domains.llm_wiki.media_structure import pdf_sections, timed_sections
from backend.domains.llm_wiki.section_assignment import normalize_assignment
from backend.domains.llm_wiki.source_structure import bind_note_structure, document_key, section_id, sections_table_id, structure_catalog
from backend.domains.vault.views.filters import apply_filter
from backend.tests.test_llm_wiki_writing_open_metadata_contract import _dependencies


def test_native_hierarchy_preserves_unnumbered_and_repeated_headings(tmp_path: Path) -> None:
    from docx import Document
    raw = ("<h1>Memory</h1><p>Main idea.</p><h2>Examples</h2><p>First example.</p>"
           "<h2>Examples</h2><p>Second example.</p><h1>Conclusion</h1><p>Final idea.</p>")
    html_segments = documents.extract_html(raw)
    paths = [segment["locator"]["section_path"] for segment in html_segments]
    assert [[node["title"] for node in path] for path in paths] == [
        ["Memory"], ["Memory", "Examples"], ["Memory", "Examples"], ["Conclusion"],
    ]
    assert paths[1][-1]["key"] != paths[2][-1]["key"]
    inserted = documents.extract_html("<h1>New opening</h1><p>New.</p>" + raw)
    assert inserted[1]["locator"]["section_path"][-1]["identity"] == paths[0][-1]["identity"]
    path = tmp_path / "source.docx"
    doc = Document()
    doc.add_heading("Memory", level=1)
    doc.add_heading("Examples", level=2)
    doc.add_paragraph("Evidence.")
    doc.save(path)
    assert [node["title"] for node in documents.extract_docx(path)[0]["locator"]["section_path"]] == ["Memory", "Examples"]
    markdown = documents.paragraph_segments("# Memory\n\nMain idea.\n\n## Examples\n\nEvidence.", locator_prefix="lines")
    assert [node["title"] for node in markdown[-1]["locator"]["section_path"]] == ["Memory", "Examples"]
    compact = documents.paragraph_segments("# Memory\nMain idea.\n## Examples\nEvidence.", locator_prefix="lines")
    assert [node["title"] for node in compact[-1]["locator"]["section_path"]] == ["Memory", "Examples"]
    assert compact[-1]["locator"]["line_start"] == 4


def test_structure_ids_survive_title_and_content_edits_and_separate_documents() -> None:
    first = {"kind": "text", "label": "book.md", "source_url": "", "content_hash": "before"}
    edited = {**first, "content_hash": "after"}
    other = {**first, "label": "other.md"}
    assert document_key(first) == document_key(edited)
    identifier = section_id("brain", "sources", "book", document_key(first), "heading-1")
    assert identifier == section_id("brain", "sources", "book", document_key(edited), "heading-1")
    assert identifier != section_id("brain", "sources", "book", document_key(other), "heading-1")
    assert identifier != section_id("brain", "sources", "other-book", document_key(first), "heading-1")
    assert document_key({"kind": "body", "label": "Old title"}) == document_key({"kind": "body", "label": "New title"})


@pytest.mark.parametrize("locale, name", [("ca", "Apartat"), ("en", "Section"), ("es", "Apartado"), ("fr", "Section")])
def test_schema_upgrade_is_idempotent_and_preserves_existing_columns(locale: str, name: str) -> None:
    custom = {"id": "custom", "name": name, "type": "text"}
    brain = {"id": "brain", "name": "Brain", "database_id": "db", "properties": [custom]}
    registry = {"tables": [brain], "views": [{"id": "main", "is_main": True,
                                               "table_id": "brain", "visibleProperties": [name]}]}
    assert ensure_source_sections(registry, brain, locale)
    assert not ensure_source_sections(registry, brain, locale)
    assert len(registry["tables"]) == 2
    assert brain["properties"][0] is custom
    relation = brain["properties"][1]
    assert relation["type"] == "relation" and relation["cardinality"] == "many-to-one"
    assert relation["relation_database_id"] == sections_table_id("brain")
    assert relation["name"] != custom["name"]
    assert registry["views"][0]["visibleProperties"] == [name, relation["name"]]


def test_timestamps_remain_locators_without_inventing_sections() -> None:
    segments = [{"text": "Early", "locator": {"start": 1, "end": 2}},
                {"text": "Later", "locator": {"start": 20, "end": 22}}]
    assert timed_sections(segments, []) == segments
    result = timed_sections(segments, [{"title": "Opening", "start_time": 0, "end_time": 10},
                                       {"title": "Discussion", "start_time": 10, "end_time": 30}])
    assert [s["locator"]["section"] for s in result] == ["Opening", "Discussion"]
    assert result[1]["locator"]["start"] == 20 and result[1]["locator"]["end"] == 22


def test_same_named_attachments_do_not_share_structure_identity(tmp_path: Path) -> None:
    from backend.services.llm_wiki_extractors import extract_resource_sources
    for folder, content in [("first", "An initial idea."), ("second", "A different idea.")]:
        path = tmp_path / folder / "book.md"
        path.parent.mkdir()
        path.write_text("# Chapter\n\n" + content)
    sources, warnings = extract_resource_sources({"Files": ["first/book.md", "second/book.md"]}, "", tmp_path,
        {"properties": [{"id": "files", "name": "Files", "type": "files"}]},
        {"attachment_property_ids": ["files"]})
    assert warnings == [] and len(sources) == 2
    assert document_key(sources[0]) != document_key(sources[1])
    note = {"source_segment_id": sources[0]["segments"][-1]["id"], "source_document_key": "old"}
    bind_note_structure([note], sources)
    assert note["source_document_key"] == document_key(sources[0])
    assert note["source_section_path"][-1]["title"] == "Chapter"


def test_pdf_uses_native_bookmarks_and_leaves_ambiguous_pages_unassigned(tmp_path: Path) -> None:
    from pypdf import PdfReader, PdfWriter
    path = tmp_path / "book.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    writer.add_blank_page(width=600, height=800)
    writer.add_outline_item("Opening", 0)
    writer.add_outline_item("Conclusion", 1)
    writer.write(path)
    assert pdf_sections(PdfReader(path))[1]["section"] == "Opening"
    writer.add_outline_item("Another heading", 1)
    writer.write(path)
    assert pdf_sections(PdfReader(path))[2] == {}
    writer = PdfWriter()
    writer.add_blank_page(width=600, height=800)
    part = writer.add_outline_item("Part", 0)
    writer.add_outline_item("Chapter", 0, parent=part)
    writer.write(path)
    assert [node["title"] for node in pdf_sections(PdfReader(path))[1]["section_path"]] == ["Part", "Chapter"]


def test_primary_segment_owns_section_even_when_context_citation_is_elsewhere() -> None:
    origin = origins.finalize_origin({"kind": "text", "label": "book.md", "input_order": 0,
        "segments": documents.paragraph_segments("# Opening\n\nEarlier evidence.\n\n# Ending\n\nFinal evidence.", locator_prefix="lines")})
    primary = origin["segments"][-1]
    context = origin["segments"][1]
    chunk = {"segments": [primary], "evidence_segments": [context]}
    note = {"title": "Final idea", "body_md": "Final idea.", "source_segment_id": primary["id"],
            "source_section_path": [{"key": "invented", "title": "Invented"}],
            "citations": [{"segment_id": context["id"], "quote": context["text"]},
                          {"segment_id": primary["id"], "quote": primary["text"]}]}
    notes, warnings = planning.validate_and_reduce_plans([(chunk, {"notes": [note]})], [origin], [],
        normalized_text=lambda value: str(value).lower(), validate_dimensions=lambda raw, specs: {})
    assert warnings == [] and notes[0]["position"] == 1
    assert notes[0]["source_section_path"][-1]["title"] == "Ending"
    assert [row["source_section_path"][-1]["title"] for row in structure_catalog([origin])] == ["Opening", "Ending"]


def test_persistence_creates_hierarchy_and_reprocessing_updates_without_duplicates(tmp_path: Path) -> None:
    brain = {"id": "brain", "name": "Brain", "properties": []}
    registry = {"tables": [brain]}
    ensure_source_sections(registry, brain, "ca")
    tables = {table["id"]: table for table in registry["tables"]}
    prop = brain["properties"][0]
    saved: dict[str, SimpleNamespace] = {}

    def save(path, metadata, body):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps({"metadata": metadata, "body": body}))
        saved[metadata["id"]] = SimpleNamespace(id=metadata["id"], path=path, metadata=dict(metadata))

    dependencies = replace(_dependencies(tmp_path), table_by_id=lambda identifier: tables[identifier],
        resolve_table_folder=lambda metadata: tmp_path / metadata["table_id"],
        save_page_md=save, parse_frontmatter=lambda raw, path: (json.loads(raw)["metadata"], json.loads(raw)["body"]),
        get_pages_for_table=lambda table_id: [page for page in saved.values() if page.metadata["table_id"] == table_id],
        page_metadata=lambda page: page.metadata, page_path=lambda page: page.path,
        uuid_factory=lambda: "note-id", merge_page_metadata=lambda metadata, page_id: metadata)
    note = {"title": "Reading idea", "body_md": "Evidence interpretation", "managed_key": "key", "position": 3,
            "source_document_key": "doc", "source_section_path": [
                {"key": "heading-1", "title": "Memory", "order": 1},
                {"key": "heading-2", "title": "Examples", "order": 2}]}
    writing.apply_plan({"notes": [note]}, "book", "Book", "brain", source_table_id="sources", dependencies=dependencies)
    sections = [page for page in saved.values() if page.metadata["table_id"] != "brain"]
    assert len(sections) == 2
    metadata = saved["note-id"].metadata
    assert len(metadata["Apartat"]) == 1
    assert len(metadata["llm_wiki_section_ancestor_ids"]) == 1
    identifier = metadata["Apartat"][0]
    note["source_section_path"][-1]["title"] = "Worked examples"
    writing.apply_plan({"notes": [note]}, "book", "Book", "brain", source_table_id="sources", dependencies=dependencies)
    assert len(saved) == 3 and saved["note-id"].metadata["Apartat"] == [identifier]
    assert saved[identifier].metadata["llm_wiki_section_path"] == "Memory › Worked examples"
    assert prop["config"]["source_sections"] is True
    note["source_section_path"] = note["source_section_path"][:1]
    writing.apply_plan({"notes": [note]}, "book", "Book", "brain", source_table_id="sources", dependencies=dependencies)
    assert saved[identifier].metadata["llm_wiki_section_stale"] is True


def test_assignment_and_hierarchy_filters_are_source_scoped() -> None:
    prop = {"name": "Apartat"}
    metadata = {"Apartat": ["child"], "llm_wiki_resource_id": "book", "llm_wiki_source_table_id": "sources"}
    row = {"id": "child", "llm_wiki_resource_id": "book", "llm_wiki_source_table_id": "sources",
           "llm_wiki_section_ancestor_ids": ["parent"]}
    normalize_assignment(metadata, prop, [row])
    assert apply_filter(metadata, "note", {"field": "Apartat", "operator": "contains", "value": "parent"})
    assert not apply_filter(metadata, "note", {"field": "Apartat", "operator": "equals", "value": "parent"})
    assert not apply_filter(metadata, "note", {"field": "Apartat", "operator": "not_contains", "value": "parent"})
    with pytest.raises(ValueError, match="only one"):
        normalize_assignment({**metadata, "Apartat": ["child", "parent"]}, prop, [row])
    with pytest.raises(ValueError, match="belong"):
        normalize_assignment({**metadata, "llm_wiki_resource_id": "other-book"}, prop, [row])
    with pytest.raises(ValueError, match="belong"):
        normalize_assignment(dict(metadata), prop, [{**row, "llm_wiki_section_stale": True}])
    normalize_assignment(dict(metadata), prop, [{**row, "llm_wiki_section_stale": True}], existing_section_id="child")
    normalize_assignment({**metadata, "Apartat": []}, prop, [row])


def test_full_schema_upgrade_preserves_a_same_named_custom_field(monkeypatch: pytest.MonkeyPatch) -> None:
    from contextlib import nullcontext
    from backend.api import vault_routes as facade
    from backend.domains.vault.knowledge import schema_service
    custom = {"id": "custom", "name": "Apartat", "type": "text"}
    brain = {"id": "brain", "properties": [custom]}
    registry = {"tables": [brain], "views": []}
    monkeypatch.setattr(facade, "registry_mutation", nullcontext)
    monkeypatch.setattr(facade, "load_registry", lambda: registry)
    monkeypatch.setattr(facade, "save_registry", lambda value: None)
    assert schema_service.ensure_brain_table_schema("brain", "ca") == 9
    assert schema_service.ensure_brain_table_schema("brain", "ca") == 0
    assert custom == {"id": "custom", "name": "Apartat", "type": "text"}
    relation = next(p for p in brain["properties"] if p.get("config", {}).get("source_sections"))
    assert relation["name"] != custom["name"]
    assert schema_service._infer_brain_roles(brain)["section"] == relation["id"]

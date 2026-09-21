"""Architecture contracts for the modular LLM Wiki extraction pipeline."""

from __future__ import annotations

from pathlib import Path

from backend.domains.llm_wiki import documents, origins
from backend.services import llm_wiki_extractors


def test_extraction_facade_preserves_origin_and_document_contracts(tmp_path: Path) -> None:
    text_path = tmp_path / "source.md"
    text_path.write_text("First paragraph.\n\nSecond paragraph.", encoding="utf-8")
    raw_origin = {
        "kind": "text",
        "label": "source.md",
        "source_url": "",
        "input_order": 0,
        "segments": documents.extract_text_file(text_path),
    }

    direct = origins.finalize_origin(raw_origin)
    compatible = llm_wiki_extractors._finalize_origin(raw_origin)  # noqa: SLF001

    assert compatible == direct
    assert llm_wiki_extractors._extract_text_file(  # noqa: SLF001
        text_path
    ) == documents.extract_text_file(text_path)


def test_extraction_modules_respect_source_guardrails() -> None:
    paths = [
        Path(llm_wiki_extractors.__file__ or ""),
        Path(documents.__file__ or ""),
        Path(origins.__file__ or ""),
    ]
    assert all(len(path.read_text(encoding="utf-8").splitlines()) <= 800 for path in paths)
    domain_sources = "\n".join(path.read_text(encoding="utf-8") for path in paths[1:])
    assert "backend.services.llm_wiki_extractors" not in domain_sources


def test_tables_and_nested_html_blocks_are_read_once_in_document_order(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "table.docx"
    document = Document()
    document.add_heading("Results", level=1)
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Measured value"
    table.cell(0, 1).text = "42"
    document.add_paragraph("After the table.")
    document.save(path)
    segments = documents.extract_docx(path)
    assert [s["text"] for s in segments] == [
        "Before the table.", "Measured value | 42", "After the table.",
    ]
    assert all(s["locator"]["section"] == "Results" for s in segments)
    assert segments[1]["locator"]["block"] == "table"
    html_segments = documents.extract_html(
        "<h1>Argument</h1><blockquote><p>Quoted position.</p></blockquote>"
        "<table><tr><td><p>Table evidence.</p></td></tr></table><p>Conclusion.</p>"
    )
    assert [s["text"] for s in html_segments] == [
        "Quoted position.", "Table evidence.", "Conclusion.",
    ]


def test_epub_uses_reading_order_instead_of_archive_insertion_order(tmp_path: Path) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("ordered-book")
    book.set_title("Ordered book")
    book.set_language("en")
    first = epub.EpubHtml(title="Opening", file_name="first.xhtml")
    first.content = "<h1>Opening</h1><p>An initial hypothesis.</p>"
    second = epub.EpubHtml(title="Conclusion", file_name="second.xhtml")
    second.content = "<h1>Conclusion</h1><p>The hypothesis is rejected.</p>"
    book.add_item(second)
    book.add_item(first)
    book.spine = [first, second]
    book.add_item(epub.EpubNcx())
    path = tmp_path / "ordered.epub"
    epub.write_epub(path, book)
    assert [s["text"] for s in documents.extract_epub(path)] == [
        "An initial hypothesis.", "The hypothesis is rejected.",
    ]
